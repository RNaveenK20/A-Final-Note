from pathlib import Path
from threading import Lock

from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook

from scripts.initialize_excel import initialize_workbook

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = BASE_DIR / "data" / "FarewellData.xlsx"
SHEET_NAME = "FarewellData"
KNOWN_NAMES = {"Praveen Athanki", "Karthikeyen Mani", "Ashok Kumar Babu", "Subbu", "Sathish", "Naveen M", "Vinod Chanda", "Aparna", "Sneha", "Ganesh", "Arul Selvan", "Padmini", "Bhuvi", "Sathya", "Uma", "Nirmal", "Kathir", "Kanaga", "Niranjana", "Raj Vishnu", "Aswathi", "Pavithran", "Manohar", "Sai", "Srinivas", "Tharun M QA", "Tharun M", "Sanjay", "Srivathsan", "Gokul", "Sayanthika"}
excel_lock = Lock()
liked_visitors = set()

app = Flask(__name__)


def clean_text(value, maximum):
    return value.strip() if isinstance(value, str) and len(value.strip()) <= maximum else None


def open_sheet():
    initialize_workbook(WORKBOOK_PATH)
    workbook = load_workbook(WORKBOOK_PATH)
    return workbook, workbook[SHEET_NAME]


def interactions():
    result = {name: {"name": name, "like": 0, "comments": []} for name in KNOWN_NAMES}
    with excel_lock:
        workbook, sheet = open_sheet()
        for name, like, comment in sheet.iter_rows(min_row=2, values_only=True):
            if name in result:
                result[name]["like"] = int(like or 0)
                if comment:
                    commenter, text = str(comment).split("\n", 1) if "\n" in str(comment) else ("A colleague", str(comment))
                    result[name]["comments"].append({"name": commenter, "text": text})
        workbook.close()
    return [result[name] for name in sorted(result)]


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/data")
def get_data():
    return jsonify(interactions())


@app.post("/like")
def like():
    payload = request.get_json(silent=True) or {}
    name = clean_text(payload.get("name"), 100)
    visitor_id = clean_text(payload.get("visitorId"), 100)
    if name not in KNOWN_NAMES:
        return jsonify(success=False, error="Unknown colleague."), 400
    if visitor_id and (visitor_id, name) in liked_visitors:
        return jsonify(success=False, error="You have already liked this message."), 409
    with excel_lock:
        workbook, sheet = open_sheet()
        rows = [row for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == name]
        if not rows:
            sheet.append([name, 0, None])
            rows = [sheet.max_row]
        count = max((int(sheet.cell(row, 2).value or 0) for row in rows), default=0) + 1
        for row in rows:
            sheet.cell(row, 2).value = count
        workbook.save(WORKBOOK_PATH)
        workbook.close()
        if visitor_id:
            liked_visitors.add((visitor_id, name))
    return jsonify(success=True, name=name, like=count)


@app.post("/comment")
def comment():
    payload = request.get_json(silent=True) or {}
    name = clean_text(payload.get("name"), 100)
    commenter = clean_text(payload.get("commenterName"), 100)
    message = clean_text(payload.get("comment"), 500)
    if name not in KNOWN_NAMES or not commenter or not message:
        return jsonify(success=False, error="Please provide a valid name and message."), 400
    with excel_lock:
        workbook, sheet = open_sheet()
        likes = [int(sheet.cell(row, 2).value or 0) for row in range(2, sheet.max_row + 1) if sheet.cell(row, 1).value == name]
        sheet.append([name, max(likes, default=0), f"{commenter}\n{message}"])
        workbook.save(WORKBOOK_PATH)
        workbook.close()
    return jsonify(success=True)


if __name__ == "__main__":
    initialize_workbook(WORKBOOK_PATH)
    app.run(debug=True)
