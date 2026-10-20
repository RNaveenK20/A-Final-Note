from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

NAMES = ["Praveen Athanki", "Karthikeyen Mani", "Ashok Kumar Babu", "Subbu", "Sathish", "Naveen M", "Vinod Chanda", "Aparna", "Sneha", "Ganesh", "Arul Selvan", "Padmini", "Bhuvi", "Sathya", "Uma", "Nirmal", "Kathir", "Kanaga", "Niranjana", "Raj Vishnu", "Aswathi", "Pavithran", "Manohar", "Sai", "Srinivas", "Tharun M QA", "Tharun M", "Sanjay", "Srivathsan", "Gokul", "Sayanthika", "Dhanusri", "Vignesh", "Tamilselvan"]


def initialize_workbook(path=None):
    path = Path(path) if path else Path(__file__).resolve().parents[1] / "data" / "FarewellData.xlsx"
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook["FarewellData"]
        existing_names = {row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]}
        for name in NAMES:
            if name not in existing_names:
                sheet.append([name, 0, None])
        sheet.auto_filter.ref = f"A1:C{sheet.max_row}"
        workbook.save(path)
        workbook.close()
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FarewellData"
    sheet.append(["Name", "Like", "Comment"])
    for name in NAMES:
        sheet.append([name, 0, None])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:C{sheet.max_row}"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 68
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


if __name__ == "__main__":
    initialize_workbook()
